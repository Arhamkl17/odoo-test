# -*- coding: utf-8 -*-
"""juni_juli_65_produk_master.py — struktur & isi kolom harga/modal ekspor produk master."""
import os

import openpyxl

DIR = "product_photos/data sheet master"
FILES = ["Produk (product.template).xlsx", "Produk (product.template) (87).xlsx",
         "mrp_bom Barang Combo.xlsx"]


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


for fn in FILES:
    path = os.path.join(DIR, fn)
    if not os.path.exists(path):
        print("!! tidak ada: %s" % path)
        continue
    print("=" * 118)
    print("BERKAS: %s" % fn)
    print("=" * 118)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        print("--- sheet '%s' : %d baris x %d kolom" % (ws.title, ws.max_row, ws.max_column))
        hdr = None
        for r in rows[:6]:
            vals = [str(c) if c is not None else "" for c in r]
            while vals and vals[-1] == "":
                vals.pop()
            print("   HDR? | %s" % " | ".join(v[:26] for v in vals))
        if not rows:
            continue
        hdr = [str(c) if c is not None else "" for c in rows[0]]
        # cari kolom yang mengandung harga/modal/biaya
        idx = {}
        for i, h in enumerate(hdr):
            hl = h.lower()
            if any(k in hl for k in ("harga", "modal", "biaya", "cost", "price", "jual", "beli")):
                idx[i] = h
        print("   kolom harga/biaya terdeteksi: %s" % idx)
        if idx:
            for i, h in idx.items():
                vals = [num(r[i]) for r in rows[1:] if i < len(r)]
                nz = [v for v in vals if v not in (None, 0)]
                print("      %-34s : %d nilai, %d != 0, min=%s max=%s" % (
                    h, len(vals), len(nz),
                    min(nz) if nz else "-", max(nz) if nz else "-"))
                if nz:
                    for r in rows[1:]:
                        if i < len(r) and num(r[i]):
                            print("          contoh: %-46s = %s" % (
                                str(r[1])[:46] if len(r) > 1 else "?", num(r[i])))
                            break
        # tampilkan 6 baris data pertama
        print("   contoh data:")
        for r in rows[1:7]:
            vals = [str(c) if c is not None else "" for c in r]
            while vals and vals[-1] == "":
                vals.pop()
            print("      | %s" % " | ".join(v[:24] for v in vals))
    wb.close()
    print("")
