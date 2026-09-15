# -*- coding: utf-8 -*-
"""
juni_juli_64_read_master_xlsx.py — DAMP IS I berkas master klien (read-only, tanpa Odoo).

  python3 scripts/juni_juli_64_read_master_xlsx.py [nama_berkas] [max_row]
"""
import os
import sys

import openpyxl

DIR = "product_photos/data sheet master"
target = sys.argv[1] if len(sys.argv) > 1 else None
MAXROW = int(sys.argv[2]) if len(sys.argv) > 2 else 200


def clean(v):
    if v is None:
        return ""
    s = str(v)
    return s if len(s) <= 34 else s[:33] + "…"


def dump(path):
    print("=" * 128)
    print("BERKAS: %s" % path)
    print("=" * 128)
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        print("  GAGAL dibuka: %r" % e)
        return
    for ws in wb.worksheets:
        print("")
        print("--- sheet '%s'  (%s baris x %s kolom)" % (ws.title, ws.max_row, ws.max_column))
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i > MAXROW:
                print("    ... (dipotong di %d baris)" % MAXROW)
                break
            cells = [clean(c) for c in row]
            while cells and cells[-1] == "":
                cells.pop()
            if not cells:
                continue
            print("   %3d | %s" % (i, " | ".join(cells)))
    wb.close()


if target:
    dump(os.path.join(DIR, target))
else:
    for fn in sorted(os.listdir(DIR)):
        if fn.lower().endswith((".xlsx", ".xls")):
            dump(os.path.join(DIR, fn))
