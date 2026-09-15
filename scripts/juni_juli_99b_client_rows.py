# -*- coding: utf-8 -*-
"""
juni_juli_99b_client_rows.py — baca baris lengkap 12 produk kembar dari ekspor produk
asli klien (`Produk (product.template).xlsx`) dan dari `05_products_with_id.csv`.

Tujuan: memutuskan apakah nama pendek & panjang itu DUA PRODUK BERBEDA (punya ID
Eksternal sendiri, harga sendiri) atau satu produk yang tercatat dua kali.
"""
import csv, os
import openpyxl

SHORT = {
    "PKG SAMBAL KOREK SURABAYA",
    "PKG SAMBAL IJO PADANG",
    "PKG SAMBAL RICA MANADO",
    "PAKET MEVVAH BERDUA",
    "PAKET KULIT CRISPY",
    "PKG INDOMIE GEPREK SAMBAL LOKAL",
}
LONG = {
    "PKG SAMBAL KOREK SURABAYA (GEPREK SAMBAL +NASI)",
    "PKG SAMBAL IJO PADANG (GEPREK SAMBAL +NASI)",
    "PKG SAMBAL RICA MANADO (GEPREK SAMBAL +NASI)",
    "PKG MEVVAH (AYAM GEPREK+TELUR ORAK ARIK+INDOMIE MEVVAH)",
    "PAKET KULIT CRISPY + NASI + MINUM",
    "PKG MIE (AYAM GEPREK+SAMBAL LOKAL+INDOMIE)",
}
ALL = SHORT | LONG
say = lambda m="": print(m)

say("=" * 126)
say("BARIS LENGKAP PRODUK KEMBAR DARI EKSPOR ASLI KLIEN")
say("=" * 126)

say("")
say("[1] Produk (product.template).xlsx  — ekspor product.template klien")
say("")
wb = openpyxl.load_workbook("product_photos/data sheet master/Produk (product.template).xlsx",
                            read_only=True, data_only=True)
for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    header = [str(h or "").strip() for h in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    say("   sheet %-16s %d kolom, %d baris" % (ws.title, len(header), len(rows) - 1))
    say("   kolom: %s" % ", ".join(header))
    say("")
    hitn = 0
    for r in rows[1:]:
        nm = str(r[idx.get("Nama", 1)] or "").strip()
        if nm in ALL:
            hitn += 1
            say("   ── %s" % nm)
            for h in header:
                i = idx[h]
                v = r[i] if i < len(r) else None
                if v not in (None, "", False):
                    say("        %-28s %s" % (h, v))
            say("")
    if not hitn:
        say("   (tidak ada baris cocok)")
    say("")
wb.close()

say("[2] Produk (product.template) (87).xlsx — ekspor klien lain (dugaan bisnis berbeda)")
say("")
try:
    wb = openpyxl.load_workbook("product_photos/data sheet master/Produk (product.template) (87).xlsx",
                                read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h or "").strip() for h in rows[0]]
        idx = {h: i for i, h in enumerate(header)}
        hit = [str(r[idx.get("Nama", 1)] or "").strip() for r in rows[1:]
               if str(r[idx.get("Nama", 1)] or "").strip() in ALL]
        say("   sheet %-16s %d baris | cocok: %d %s" % (ws.title, len(rows) - 1, len(hit), hit[:4]))
    wb.close()
except Exception as e:
    say("   GAGAL: %s" % e)

say("")
say("[3] 05_products_with_id.csv — berkas yang dipakai impor")
say("")
with open("import_data/csv/05_products_with_id.csv", encoding="utf-8-sig") as f:
    rd = csv.DictReader(f)
    for row in rd:
        nm = (row.get("Nama") or "").strip()
        if nm in ALL:
            say("   %-62s | harga jual=%-10s modal=%-8s unit=%-8s id=%s" % (
                nm[:62], row.get("Harga Jual"), row.get("Modal"),
                row.get("Unit Pembelian"), row.get("ID Eksternal")))

say("")
say("[4] Jumlah total produk di masing-masing sumber")
say("")
with open("import_data/csv/05_products_with_id.csv", encoding="utf-8-sig") as f:
    n5 = sum(1 for _ in csv.DictReader(f))
say("   05_products_with_id.csv                 : %d baris" % n5)
wb = openpyxl.load_workbook("product_photos/data sheet master/Produk (product.template).xlsx",
                            read_only=True, data_only=True)
for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    say("   Produk (product.template).xlsx [%s]     : %d baris" % (ws.title, max(len(rows) - 1, 0)))
wb.close()
say("")
say("=" * 126)
