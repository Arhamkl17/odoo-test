# -*- coding: utf-8 -*-
"""
juni_juli_98_source_names.py — cari nama 12 produk kembar di SEMUA berkas sumber klien.

Tujuan: memutuskan mana yang asli, dengan bukti nama itu benar-benar ada di berkas klien.
Tidak menyentuh database sama sekali (hanya baca berkas).
"""
import csv, glob, os, re

TARGET = [
    "PKG SAMBAL KOREK SURABAYA",
    "PKG SAMBAL KOREK SURABAYA (GEPREK SAMBAL +NASI)",
    "PKG SAMBAL IJO PADANG",
    "PKG SAMBAL IJO PADANG (GEPREK SAMBAL +NASI)",
    "PKG SAMBAL RICA MANADO",
    "PKG SAMBAL RICA MANADO (GEPREK SAMBAL +NASI)",
    "PAKET MEVVAH BERDUA",
    "PKG MEVVAH (AYAM GEPREK+TELUR ORAK ARIK+INDOMIE MEVVAH)",
    "PAKET KULIT CRISPY",
    "PAKET KULIT CRISPY + NASI + MINUM",
    "PKG INDOMIE GEPREK SAMBAL LOKAL",
    "PKG MIE (AYAM GEPREK+SAMBAL LOKAL+INDOMIE)",
]

say = lambda m="": print(m)
say("=" * 118)
say("PENCARIAN NAMA DI BERKAS SUMBER KLIEN")
say("=" * 118)

# ---------------------------------------------------------------- CSV
say("")
say("[CSV] import_data/csv/*.csv")
say("")
for path in sorted(glob.glob("import_data/csv/*.csv")):
    try:
        with open(path, encoding="utf-8-sig", errors="ignore") as f:
            txt = f.read()
    except Exception as e:
        say("   %-46s GAGAL: %s" % (os.path.basename(path), e))
        continue
    hits = [t for t in TARGET if t in txt]
    if hits:
        say("   %-46s" % os.path.basename(path))
        for h in hits:
            say("        ✓ %s" % h)

# ---------------------------------------------------------------- JSON
say("")
say("[JSON] import_data/*.json")
say("")
for path in sorted(glob.glob("import_data/*.json")):
    try:
        with open(path, encoding="utf-8", errors="ignore") as f:
            txt = f.read()
    except Exception as e:
        say("   %-46s GAGAL: %s" % (os.path.basename(path), e))
        continue
    hits = sorted({t for t in TARGET if t in txt})
    say("   %-46s %s" % (os.path.basename(path), ("%d nama cocok" % len(hits)) if hits else "—"))
    for h in hits:
        say("        ✓ %s" % h)

# ---------------------------------------------------------------- XLSX master
say("")
say("[XLSX] product_photos/data sheet master/*.xlsx  &  berkas master lain")
say("")
try:
    import openpyxl
except ImportError:
    say("   openpyxl tidak tersedia")
    openpyxl = None

MASTER = sorted(glob.glob("product_photos/data sheet master/*.xlsx")) + \
    sorted(glob.glob("import_data/*.xlsx")) + sorted(glob.glob("*.xlsx"))

for path in MASTER:
    if openpyxl is None:
        break
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        say("   %-60s GAGAL: %s" % (os.path.basename(path), e))
        continue
    found = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str):
                    s = cell.strip()
                    if s in TARGET:
                        found.setdefault(s, set()).add(ws.title)
    say("   %s" % os.path.relpath(path))
    if found:
        for nama, sheets in sorted(found.items()):
            say("        ✓ %-62s sheet: %s" % (nama[:62], ", ".join(sorted(sheets))))
    else:
        say("        —")
    wb.close()

# ---------------------------------------------------------------- ringkasan
say("")
say("=" * 118)
say("RINGKASAN — di berapa berkas masing-masing nama ditemukan")
say("")
found_in = {t: [] for t in TARGET}
for path in sorted(glob.glob("import_data/csv/*.csv")) + sorted(glob.glob("import_data/*.json")):
    try:
        txt = open(path, encoding="utf-8-sig", errors="ignore").read()
    except Exception:
        continue
    for t in TARGET:
        if t in txt:
            found_in[t].append(os.path.basename(path))
if openpyxl:
    for path in MASTER:
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        hits = set()
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if isinstance(cell, str) and cell.strip() in TARGET:
                        hits.add(cell.strip())
        for h in hits:
            found_in[h].append(os.path.basename(path))
        wb.close()

for t in TARGET:
    files = sorted(set(found_in[t]))
    say("   %-62s %s" % (t[:62], ("%d berkas: %s" % (len(files), ", ".join(files[:3]))) if files else "TIDAK ADA di berkas mana pun"))
say("")
say("=" * 118)
